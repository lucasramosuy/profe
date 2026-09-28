"""Hoja A4 uniforme para material editable; nunca refluye PDF ni presentaciones."""
from pathlib import Path
import re
import unicodedata
from xml.etree import ElementTree as ET

from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

INK = HexColor('#211d1a')
MUTED = HexColor('#6a625b')
CORAL = HexColor('#bb5943')
RULE = HexColor('#d6d1c9')
PAPER = HexColor('#fffefa')
W, H = A4
LEFT, RIGHT, TOP, BOTTOM = 49, 49, 84, 58
WIDTH = W - LEFT - RIGHT
FONTS = Path('/usr/share/fonts/truetype/dejavu')


def fonts():
    for face, filename in (('ProfeSans', 'DejaVuSans.ttf'), ('ProfeBold', 'DejaVuSans-Bold.ttf')):
        font = FONTS / filename
        if not font.is_file():
            raise RuntimeError('Falta DejaVu Sans en el runner; no se generó el imprimible')
        pdfmetrics.registerFont(TTFont(face, str(font)))


def clean(text):
    return ''.join(ch for ch in str(text) if ch in '\n\t' or unicodedata.category(ch)[0] != 'C').strip()


def from_docx(path):
    from docx import Document
    doc = Document(path)
    if doc.tables or any(p._element.xpath('.//w:drawing|.//w:pict|.//w:object|.//w:fldSimple') for p in doc.paragraphs):
        raise ValueError('El DOCX tiene tablas, gráficos o campos: conservá su formato original en PDF')
    blocks = []
    for para in doc.paragraphs:
        text = clean(para.text)
        if not text:
            continue
        style = para.style.name.lower() if para.style else ''
        if 'heading' in style or 'título' in style or 'titulo' in style:
            kind = 'heading'
        elif 'list' in style or para._element.xpath('./w:pPr/w:numPr'):
            kind = 'bullet'
        else:
            kind = 'paragraph'
        blocks.append((kind, text))
    return blocks


def from_odt(path):
    from zipfile import ZipFile
    with ZipFile(path) as z:
        xml = ET.fromstring(z.read('content.xml'))
    ns = {'text': 'urn:oasis:names:tc:opendocument:xmlns:text:1.0',
          'table': 'urn:oasis:names:tc:opendocument:xmlns:table:1.0',
          'draw': 'urn:oasis:names:tc:opendocument:xmlns:drawing:1.0'}
    if xml.find('.//table:table', ns) is not None or xml.find('.//draw:frame', ns) is not None:
        raise ValueError('El ODT tiene tablas o imágenes: conservá su formato original en PDF')
    blocks = []
    for node in xml.iter():
        tag = node.tag.split('}')[-1]
        if tag in ('h', 'p') and node.tag.startswith('{' + ns['text'] + '}'):
            text = clean(''.join(node.itertext()))
            if text:
                blocks.append(('heading' if tag == 'h' else 'paragraph', text))
    return blocks


def from_text(path):
    text = path.read_text(encoding='utf-8-sig')
    if len(text) > 200_000:
        raise ValueError('El texto supera 200 KB; dividilo en materiales más cortos')
    blocks = []
    for line in text.splitlines():
        line = clean(line)
        if not line:
            continue
        if path.suffix.lower() == '.md':
            if re.match(r'^#{1,6}\s', line):
                blocks.append(('heading', re.sub(r'^#{1,6}\s+', '', line)))
                continue
            if re.match(r'^[-*+]\s', line):
                blocks.append(('bullet', re.sub(r'^[-*+]\s+', '', line)))
                continue
            if re.match(r'^\d+[.)]\s', line):
                blocks.append(('bullet', line))
                continue
            if any(marker in line for marker in ('![', '|', '```', '<img')):
                raise ValueError('El Markdown tiene medios, tabla o código: no se remaqueta automáticamente')
        blocks.append(('paragraph', line))
    return blocks


def wrap(text, face, size, width):
    lines, current = [], ''
    for word in text.split():
        candidate = (current + ' ' + word).strip()
        if pdfmetrics.stringWidth(candidate, face, size) > width and current:
            lines.append(current)
            current = word
        else:
            current = candidate
        if pdfmetrics.stringWidth(current, face, size) > width:
            raise ValueError('Hay una palabra o URL demasiado larga para maquetar sin cortar')
    if current:
        lines.append(current)
    return lines


def render(src, target, title, course, description):
    fonts()
    ext = src.suffix.lower()
    if ext == '.docx':
        blocks = from_docx(src)
    elif ext == '.odt':
        blocks = from_odt(src)
    elif ext in ('.txt', '.md'):
        blocks = from_text(src)
    else:
        raise ValueError('Solo DOCX, ODT, TXT y Markdown admiten plantilla')
    if not blocks:
        raise ValueError('El material no tiene texto para maquetar')
    if len(blocks) > 1500:
        raise ValueError('El material es demasiado largo para un imprimible')
    title = clean(title)
    course = clean(course)
    description = clean(description)
    if not title or len(title) > 180 or len(course) > 100:
        raise ValueError('Revisá título y curso antes de imprimir')
    c = canvas.Canvas(str(target), pagesize=A4)
    page = 0
    y = 0

    def start():
        nonlocal page, y
        if page:
            c.showPage()
        page += 1
        if page > 40:
            raise ValueError('El material supera 40 páginas; dividilo antes de imprimir')
        c.setFillColor(PAPER); c.rect(0, 0, W, H, fill=1, stroke=0)
        c.setFillColor(CORAL); c.rect(0, H-208, 9, 208, stroke=0, fill=1)
        c.setFillColor(MUTED); c.setFont('ProfeSans', 8)
        c.drawRightString(W-RIGHT, H-51, 'MATERIAL DE CLASE' + ('  /  ' + course if course else ''))
        c.drawRightString(W-RIGHT, 24, f'{page:02d}')
        y = H-TOP

    def space(height):
        nonlocal y
        if y-height < BOTTOM:
            start()

    def block(kind, text):
        nonlocal y
        if kind == 'heading':
            face, size, leading, before, after = 'ProfeBold', 13, 20, 12, 10
        elif kind == 'bullet':
            face, size, leading, before, after = 'ProfeSans', 10, 17, 2, 8
        else:
            face, size, leading, before, after = 'ProfeSans', 10, 17, 0, 9
        prefix = '•  ' if kind == 'bullet' and not re.match(r'^\d+[.)]', text) else ''
        lines = wrap(prefix+text, face, size, WIDTH)
        height = before + len(lines)*leading + after
        if height > H-TOP-BOTTOM:
            raise ValueError('Un bloque es demasiado largo para la página')
        space(height)
        y -= before
        c.setFillColor(INK if kind == 'heading' else MUTED if kind == 'bullet' else INK)
        c.setFont(face, size)
        for line in lines:
            c.drawString(LEFT, y, line)
            y -= leading
        y -= after

    start()
    # Same v3 hierarchy and whitespace, without the wordmark or a fixed diagram.
    c.setFillColor(CORAL); c.setFont('ProfeBold', 8.5)
    c.drawString(LEFT, y-23, 'MATERIAL  /  IMPRIMIBLE')
    y -= 53
    for line in wrap(title, 'ProfeBold', 22, WIDTH):
        space(30)
        c.setFillColor(INK); c.setFont('ProfeBold', 22)
        c.drawString(LEFT, y, line)
        y -= 30
    y -= 10
    if description:
        block('paragraph', description)
    y -= 14
    for kind, text in blocks:
        block(kind, text)
    c.setTitle(title)
    c.save()
    return page
